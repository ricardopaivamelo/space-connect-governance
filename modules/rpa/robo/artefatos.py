"""
artefatos.py
============
Etapa final do fluxo: a partir dos resultados consolidados, gera
automaticamente os artefatos técnicos (critério "Entrega de Artefatos"):

  - Planilha Excel estruturada (uma aba por tipo + aba de resumo com fórmulas)
  - Relatório PDF consolidado
  - JSON estruturado (consumido pelo dashboard / front-end)
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
)

from .config import Config

AZUL = "0B3D91"
AZUL_HEX = colors.HexColor("#0B3D91")
VERMELHO = colors.HexColor("#C62828")
LARANJA = colors.HexColor("#EF6C00")
VERDE = colors.HexColor("#2E7D32")

COR_SEVERIDADE = {
    "CRÍTICO": "FFCDD2", "ANOMALIA": "FFCDD2",
    "ALERTA": "FFE0B2",
    "NOMINAL": "C8E6C9",
}


# ===========================================================================
# EXCEL
# ===========================================================================
def _estilizar_cabecalho(ws, n_colunas: int):
    fill = PatternFill("solid", start_color=AZUL)
    fonte = Font(bold=True, color="FFFFFF", name="Arial")
    for col in range(1, n_colunas + 1):
        c = ws.cell(row=1, column=col)
        c.fill = fill
        c.font = fonte
        c.alignment = Alignment(horizontal="center", vertical="center")


def _auto_largura(ws):
    for col in ws.columns:
        largura = max((len(str(c.value)) for c in col if c.value is not None),
                      default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            min(largura + 3, 50)


def _aba_severidade(ws, coluna_idx: int, n_linhas: int):
    """Pinta a célula de severidade conforme o nível."""
    for r in range(2, n_linhas + 2):
        c = ws.cell(row=r, column=coluna_idx)
        cor = COR_SEVERIDADE.get(str(c.value), None)
        if cor:
            c.fill = PatternFill("solid", start_color=cor)
            c.font = Font(bold=True, name="Arial")


def gerar_excel(resultados: Dict[str, List[Dict]], cfg: Config) -> Path:
    wb = Workbook()

    # ---- Aba RESUMO ----
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["Tipo de dado", "Itens processados", "Itens com alerta/anomalia"])
    _estilizar_cabecalho(ws, 3)

    tel, doc, img = (resultados["telemetria"], resultados["documentos"],
                     resultados["imagens"])

    def _conta_criticos(lista, niveis):
        return sum(1 for x in lista if x.get("severidade") in niveis)

    linhas_resumo = [
        ("Telemetria", len(tel), _conta_criticos(tel, {"CRÍTICO"})),
        ("Documentos", len(doc), _conta_criticos(doc, {"CRÍTICO", "ALERTA"})),
        ("Imagens", len(img), _conta_criticos(img, {"ANOMALIA"})),
    ]
    for nome, total, criticos in linhas_resumo:
        ws.append([nome, total, criticos])

    # Linha de TOTAL usando FÓRMULAS (não valores fixos)
    ws.append(["TOTAL", "=SUM(B2:B4)", "=SUM(C2:C4)"])
    for col in (1, 2, 3):
        ws.cell(row=5, column=col).font = Font(bold=True, name="Arial")
    _auto_largura(ws)

    # ---- Aba TELEMETRIA ----
    if tel:
        ws = wb.create_sheet("Telemetria")
        ws.append(["Arquivo", "Linhas", "Variáveis monitoradas",
                   "Outliers (z-score)", "Flags ML (IsolationForest)",
                   "Severidade", "Resumo"])
        _estilizar_cabecalho(ws, 7)
        for r in tel:
            ws.append([
                r["nome"], r["linhas"], len(r["colunas_monitoradas"]),
                r["n_anomalias_zscore"], r["n_anomalias_isolationforest"],
                r["severidade"], r["resumo"],
            ])
        _aba_severidade(ws, 6, len(tel))
        _auto_largura(ws)

    # ---- Aba DOCUMENTOS ----
    if doc:
        ws = wb.create_sheet("Documentos")
        ws.append(["Arquivo", "Palavras", "Severidade", "Códigos de erro",
                   "Subsistemas", "Datas", "Resumo"])
        _estilizar_cabecalho(ws, 7)
        for r in doc:
            e = r["entidades"]
            ws.append([
                r["nome"], r["palavras"], r["severidade"],
                ", ".join(e["codigos_erro"]) or "-",
                ", ".join(e["subsistemas"]) or "-",
                ", ".join(e["datas"]) or "-",
                r["resumo"],
            ])
        _aba_severidade(ws, 3, len(doc))
        _auto_largura(ws)

    # ---- Aba IMAGENS ----
    if img:
        ws = wb.create_sheet("Imagens")
        ws.append(["Arquivo", "Dimensões", "Brilho médio", "Contraste",
                   "Nitidez", "% Hotspots", "Severidade", "Resumo"])
        _estilizar_cabecalho(ws, 8)
        for r in img:
            ws.append([
                r["nome"], f"{r['largura']}x{r['altura']}", r["brilho_medio"],
                r["contraste"], r["nitidez"], r["perc_hotspots"],
                r["severidade"], r["resumo"],
            ])
        _aba_severidade(ws, 7, len(img))
        _auto_largura(ws)

    caminho = cfg.pasta_outputs / cfg.arq_excel
    wb.save(caminho)
    return caminho


# ===========================================================================
# JSON (para o dashboard / front-end)
# ===========================================================================
def gerar_json(resultados: Dict, estatisticas: Dict, cfg: Config) -> Path:
    payload = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "estatisticas": estatisticas,
        "resultados": resultados,
    }
    caminho = cfg.pasta_outputs / cfg.arq_json
    caminho.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                       encoding="utf-8")
    return caminho


# ===========================================================================
# PDF (relatório consolidado)
# ===========================================================================
def gerar_pdf(resultados: Dict, estatisticas: Dict, cfg: Config) -> Path:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("Tit", parent=styles["Title"], textColor=AZUL_HEX,
                              fontSize=20))
    styles.add(ParagraphStyle("Sub", parent=styles["Normal"],
                              alignment=TA_CENTER, fontSize=11,
                              textColor=colors.HexColor("#37474F")))
    styles.add(ParagraphStyle("H", parent=styles["Heading2"],
                              textColor=AZUL_HEX, fontSize=13, spaceBefore=12))
    styles.add(ParagraphStyle("P", parent=styles["Normal"], fontSize=10,
                              leading=14))

    styles.add(ParagraphStyle("Cel", parent=styles["Normal"], fontSize=8.5,
                              leading=11))

    el = []
    el.append(Paragraph("Relatório de Triagem de Missão", styles["Tit"]))
    el.append(Paragraph("MissionOps RPA — análise automatizada multimodal",
                        styles["Sub"]))
    el.append(Spacer(1, 0.3 * cm))
    el.append(Paragraph(
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Sub"]))
    el.append(Spacer(1, 0.4 * cm))
    el.append(HRFlowable(width="100%", color=AZUL_HEX, thickness=1))

    # ---- Painel de indicadores ----
    el.append(Paragraph("1. Indicadores gerais", styles["H"]))
    kpi = [
        ["Arquivos processados", str(estatisticas["total_processados"])],
        ["Falhas de processamento", str(estatisticas["total_falhas"])],
        ["Telemetrias com anomalia", str(estatisticas["telemetria_criticas"])],
        ["Documentos críticos/alerta", str(estatisticas["documentos_criticos"])],
        ["Imagens com anomalia", str(estatisticas["imagens_anomalas"])],
    ]
    t = Table(kpi, colWidths=[9 * cm, 4 * cm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ECEFF1")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0BEC5")),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    el.append(t)

    # ---- Itens que exigem atenção ----
    el.append(Paragraph("2. Itens que exigem atenção", styles["H"]))
    criticos = []
    for r in resultados["telemetria"]:
        if r["severidade"] == "CRÍTICO":
            criticos.append([r["nome"], "Telemetria",
                             Paragraph(r["resumo"], styles["Cel"])])
    for r in resultados["documentos"]:
        if r["severidade"] in ("CRÍTICO", "ALERTA"):
            criticos.append([r["nome"], f"Documento ({r['severidade']})",
                             Paragraph(r["resumo"], styles["Cel"])])
    for r in resultados["imagens"]:
        if r["severidade"] == "ANOMALIA":
            criticos.append([r["nome"], "Imagem",
                             Paragraph(r["resumo"], styles["Cel"])])

    if criticos:
        dados = [["Arquivo", "Tipo", "Observação"]] + criticos
        tab = Table(dados, colWidths=[4 * cm, 3.5 * cm, 9 * cm])
        tab.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), AZUL_HEX),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0BEC5")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#FFF3E0")]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        el.append(tab)
    else:
        el.append(Paragraph("Nenhum item crítico identificado nesta execução.",
                            styles["P"]))

    # ---- Falhas ----
    if estatisticas["falhas"]:
        el.append(Paragraph("3. Falhas de processamento (error handling)",
                            styles["H"]))
        for f in estatisticas["falhas"]:
            el.append(Paragraph(f"• <b>{f['arquivo']}</b>: {f['erro']}",
                                styles["P"]))

    caminho = cfg.pasta_outputs / cfg.arq_pdf
    doc = SimpleDocTemplate(str(caminho), pagesize=A4,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    doc.build(el)
    return caminho
